# -*- coding: utf-8 -*-
"""產生 NoteAdventure 期中專題簡報（白／橘風格 + 實際前端截圖）。"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parent
SHOT = ROOT / 'screenshots_crop'
if not SHOT.exists():
    SHOT = ROOT / 'screenshots'
OUT = ROOT / 'NoteAdventure_Midterm_Report.pptx'

# 與前端相同色票
ORANGE = RGBColor(0xFF, 0x5A, 0x52)
ORANGE_SOFT = RGBColor(0xFF, 0xF7, 0xF7)
ORANGE_BORDER = RGBColor(0xFF, 0xD6, 0xD3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF7, 0xF7, 0xF8)
DARK = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x37, 0x41, 0x51)
LIGHT_GRAY = RGBColor(0xD1, 0xD5, 0xDB)
MUTED = RGBColor(0x6B, 0x72, 0x80)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size=18, bold=False, color=DARK, name='Microsoft JhengHei'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()


def add_pixel_bar(slide, top=0, height=Inches(0.28)):
    # 頂部橘色像素條
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    shape.line.fill.background()
    # 右側小色塊營造像素感
    for i, color in enumerate([ORANGE, RGBColor(0xFF, 0x7A, 0x72), ORANGE_BORDER, WHITE]):
        block = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            W - Inches(0.35) * (i + 1),
            top,
            Inches(0.35),
            height
        )
        block.fill.solid()
        block.fill.fore_color.rgb = color
        block.line.fill.background()


def add_footer(slide, page_no, total):
    box = slide.shapes.add_textbox(Inches(0.5), H - Inches(0.45), Inches(10), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f'NoteAdventure ｜ 期中專題報告 ｜ {page_no}/{total}'
    set_run(run, 11, color=MUTED)
    # 右下角橘點
    dot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        W - Inches(0.55),
        H - Inches(0.45),
        Inches(0.22),
        Inches(0.22)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = ORANGE
    dot.line.fill.background()


def add_title(slide, text, left=Inches(0.6), top=Inches(0.45), width=Inches(12)):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.7))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, 32, bold=True, color=DARK)
    # 標題下橘線
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left,
        top + Inches(0.62),
        Inches(1.2),
        Inches(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.fill.background()


def add_body(slide, lines, left=Inches(0.6), top=Inches(1.4), width=Inches(12), size=18):
    box = slide.shapes.add_textbox(left, top, width, Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = line
        set_run(run, size, color=GRAY)


def add_card(slide, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LIGHT_GRAY
    card.line.width = Pt(1.5)
    return card


def add_image_fit(slide, path, left, top, width, height):
    if not path.exists():
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def screenshot_slide(prs, title, image_name, bullets, page_no, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_pixel_bar(slide)
    add_title(slide, title)
    add_card(slide, Inches(0.5), Inches(1.35), Inches(8.2), Inches(5.4))
    img = SHOT / image_name
    if img.exists():
        # 固定高度避免截圖撐破版面
        slide.shapes.add_picture(
            str(img),
            Inches(0.65),
            Inches(1.5),
            width=Inches(7.9),
            height=Inches(5.05)
        )
    add_card(slide, Inches(9.0), Inches(1.35), Inches(3.8), Inches(5.4))
    box = slide.shapes.add_textbox(Inches(9.2), Inches(1.55), Inches(3.4), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = '重點'
    set_run(run, 18, bold=True, color=ORANGE)
    for line in bullets:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = f'• {line}'
        set_run(run, 14, color=GRAY)
    add_footer(slide, page_no, total)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 15

    # 1 封面
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s, 0, Inches(0.55))
    for i in range(8):
        block = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.4) + Inches(0.35) * i,
            Inches(0.55),
            Inches(0.28),
            Inches(0.28)
        )
        block.fill.solid()
        block.fill.fore_color.rgb = ORANGE if i % 2 == 0 else ORANGE_BORDER
        block.line.fill.background()
    add_card(s, Inches(2.2), Inches(1.8), Inches(8.9), Inches(4.0))
    box = s.shapes.add_textbox(Inches(2.5), Inches(2.2), Inches(8.3), Inches(3.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'NoteAdventure'
    set_run(run, 48, bold=True, color=DARK)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = '音樂練習 RPG 管理系統'
    set_run(run, 24, color=GRAY)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = '期中專題報告 ｜ 約 10–15 分鐘'
    set_run(run, 16, color=ORANGE)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'Streamlit → Flask API → MySQL'
    set_run(run, 14, color=MUTED)
    add_footer(s, 1, total)

    # 2 大綱
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s)
    add_title(s, '報告大綱')
    add_body(s, [
        '1. 為什麼做這個系統？',
        '2. 專案目標與使用者',
        '3. 技術架構（前端／API／資料庫）',
        '4. 核心流程：任務 → 練習 → 獎勵 → 成長',
        '5. 實際操作演示（老師端／學生端截圖）',
        '6. 目前完成進度與下一步',
        '7. 結語與 Q&A',
    ], size=20)
    add_footer(s, 2, total)

    # 3 動機
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s)
    add_title(s, '為什麼需要 NoteAdventure？')
    cards = [
        ('問題', '一對一音樂老師難掌握學生練習狀況，學生也容易中斷練習。'),
        ('解法', '把練習變成可追蹤的任務，完成後給 XP、金幣與寵物成長。'),
        ('價值', '老師管理更輕鬆；學生有明確目標與回饋，提升持續練習意願。'),
    ]
    for i, (t, body) in enumerate(cards):
        left = Inches(0.6) + Inches(4.1) * i
        add_card(s, left, Inches(1.5), Inches(3.8), Inches(4.5))
        accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.5), Inches(3.8), Inches(0.18))
        accent.fill.solid()
        accent.fill.fore_color.rgb = ORANGE
        accent.line.fill.background()
        box = s.shapes.add_textbox(left + Inches(0.25), Inches(1.9), Inches(3.3), Inches(3.8))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = t
        set_run(run, 22, bold=True, color=ORANGE)
        p = tf.add_paragraph()
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = body
        set_run(run, 16, color=GRAY)
    add_footer(s, 3, total)

    # 4 目標
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s)
    add_title(s, '專案目標與角色')
    add_card(s, Inches(0.6), Inches(1.5), Inches(5.8), Inches(4.8))
    add_card(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(4.8))
    box = s.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(5.2), Inches(4.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = '老師'
    set_run(run, 24, bold=True, color=ORANGE)
    for line in [
        '註冊／登入',
        '建立學生、重設密碼',
        '建立與指派任務',
        '查看上傳紀錄',
        '管理商店與兌換',
        '首頁儀表板總覽',
    ]:
        p = tf.add_paragraph()
        run = p.add_run()
        run.text = f'• {line}'
        set_run(run, 16, color=GRAY)
    box = s.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.2), Inches(4.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = '學生'
    set_run(run, 24, bold=True, color=ORANGE)
    for line in [
        '登入與首次強制改密碼',
        '查看任務與進度',
        '上傳練習錄音',
        '獲得 XP／金幣',
        '寵物成長',
        '商店購買與我的背包',
    ]:
        p = tf.add_paragraph()
        run = p.add_run()
        run.text = f'• {line}'
        set_run(run, 16, color=GRAY)
    add_footer(s, 4, total)

    # 5 架構
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s)
    add_title(s, '技術架構')
    layers = [
        ('Streamlit 前端', '登入、操作介面\nHTTP + session cookie'),
        ('Flask API', '驗證、參數檢查\n呼叫 services'),
        ('Services + MySQL', '商業邏輯\nmusic_rpg_v2'),
    ]
    for i, (t, body) in enumerate(layers):
        left = Inches(0.8) + Inches(4.0) * i
        add_card(s, left, Inches(2.0), Inches(3.6), Inches(3.6))
        topbar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(2.0), Inches(3.6), Inches(0.7))
        topbar.fill.solid()
        topbar.fill.fore_color.rgb = ORANGE
        topbar.line.fill.background()
        box = s.shapes.add_textbox(left + Inches(0.2), Inches(2.15), Inches(3.2), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = t
        set_run(run, 18, bold=True, color=WHITE)
        box = s.shapes.add_textbox(left + Inches(0.3), Inches(3.0), Inches(3.0), Inches(2.2))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = body
        set_run(run, 16, color=GRAY)
        if i < 2:
            arrow = s.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                left + Inches(3.65),
                Inches(3.5),
                Inches(0.35),
                Inches(0.35)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ORANGE
            arrow.line.fill.background()
    add_footer(s, 5, total)

    # 6 流程
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s)
    add_title(s, '核心資料流程')
    steps = ['建立學生', '建立／指派任務', '上傳練習', '更新進度', '發放獎勵', '寵物成長']
    for i, step in enumerate(steps):
        left = Inches(0.5) + Inches(2.1) * i
        add_card(s, left, Inches(2.5), Inches(1.95), Inches(2.2))
        num = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left + Inches(0.7),
            Inches(2.75),
            Inches(0.55),
            Inches(0.55)
        )
        num.fill.solid()
        num.fill.fore_color.rgb = ORANGE
        num.line.fill.background()
        box = s.shapes.add_textbox(left + Inches(0.7), Inches(2.82), Inches(0.55), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(i + 1)
        set_run(run, 18, bold=True, color=WHITE)
        box = s.shapes.add_textbox(left + Inches(0.1), Inches(3.5), Inches(1.75), Inches(1.0))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = step
        set_run(run, 14, bold=True, color=DARK)
    add_body(s, [
        '技術重點：獎勵發放使用 transaction，避免重複發獎；購買扣款與庫存同步處理。'
    ], top=Inches(5.2), size=16)
    add_footer(s, 6, total)

    # 7-13 截圖頁
    demos = [
        ('登入與老師註冊', '01_login.png', [
            '老師／學生分頁登入',
            '支援老師註冊',
            '前端透過 Flask API 驗證',
        ]),
        ('老師首頁儀表板', '02_teacher_home.png', [
            '學生數、任務、上傳統計',
            '進行中任務一覽',
            '登入後預設進入首頁',
        ]),
        ('學生管理', '03_teacher_students.png', [
            '查看學生 XP／金幣',
            '可建立學生與重設密碼',
            '帳號由系統自動產生',
        ]),
        ('任務清單與指派', '04_teacher_quests.png', [
            '建立每日／每週／隱藏任務',
            '指派後有期限與進度',
            '完成後自動發獎',
        ]),
        ('練習上傳與商品', '05_teacher_uploads.png', [
            '老師查看學生上傳',
            '商店建立／待兌換',
            '售完自動下架',
        ]),
        ('學生首頁與寵物', '09_student_home.png', [
            'XP／金幣／任務摘要',
            '寵物階段與進度',
            '最近上傳與購買',
        ]),
        ('商店與我的背包', '13_student_backpack.png', [
            '商店購買虛擬／實體商品',
            '背包顯示名稱與說明',
            '5×6 方格背包介面',
        ]),
    ]
    page = 7
    for title, img, bullets in demos:
        screenshot_slide(prs, title, img, bullets, page, total)
        page += 1

    # 14 進度
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s)
    add_title(s, '目前完成進度')
    add_body(s, [
        '已完成：帳號、任務、上傳、進度獎勵、寵物、商店／兌換、Dashboard、我的背包、老師註冊',
        '架構：Console + Flask API + Streamlit 三層可共用 services',
        '進行中：期中展示流程完整演練',
        '可選下一步：任務／商品進階管理、更多任務條件、自動化測試',
    ], size=18)
    add_footer(s, 14, total)

    # 15 結語
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_pixel_bar(s, 0, Inches(0.55))
    add_card(s, Inches(2.0), Inches(1.8), Inches(9.3), Inches(3.8))
    box = s.shapes.add_textbox(Inches(2.4), Inches(2.2), Inches(8.5), Inches(3.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = '謝謝聆聽'
    set_run(run, 40, bold=True, color=DARK)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'NoteAdventure ｜ 把練習變成冒險'
    set_run(run, 20, color=ORANGE)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'Q & A'
    set_run(run, 24, bold=True, color=GRAY)
    add_footer(s, 15, total)

    prs.save(OUT)
    print('saved', OUT)


if __name__ == '__main__':
    build()
