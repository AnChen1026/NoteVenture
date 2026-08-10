# -*- coding: utf-8 -*-
"""產生 NoteAdventure 期中簡報（含 Console／Streamlit／學生 UI 實際畫面）。"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SHOT = ROOT / 'screenshots_crop'
if not SHOT.exists():
    SHOT = ROOT / 'screenshots'
OUT = ROOT / 'NoteAdventure_Midterm_Report.pptx'

ORANGE = RGBColor(0xFF, 0x5A, 0x52)
ORANGE_SOFT = RGBColor(0xFF, 0xF7, 0xF7)
ORANGE_BORDER = RGBColor(0xFF, 0xD6, 0xD3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF7, 0xF7, 0xF8)
DARK = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x37, 0x41, 0x51)
MUTED = RGBColor(0x6B, 0x72, 0x80)

W = Inches(13.333)
H = Inches(7.5)
TOTAL = 16


def set_run(run, size=18, bold=False, color=DARK, name='Microsoft JhengHei'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()


def add_pixel_bar(slide, top=0, height=Inches(0.28)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    shape.line.fill.background()
    for i, color in enumerate([ORANGE, RGBColor(0xFF, 0x7A, 0x72), ORANGE_BORDER, WHITE]):
        block = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            W - Inches(0.35) * (i + 1),
            top,
            Inches(0.35),
            height,
        )
        block.fill.solid()
        block.fill.fore_color.rgb = color
        block.line.fill.background()


def add_footer(slide, page_no):
    box = slide.shapes.add_textbox(Inches(0.5), H - Inches(0.45), Inches(10), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f'NoteAdventure ｜ 期中專題報告 ｜ {page_no}/{TOTAL}'
    set_run(run, 11, color=MUTED)
    dot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        W - Inches(0.55),
        H - Inches(0.45),
        Inches(0.22),
        Inches(0.22),
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = ORANGE
    dot.line.fill.background()


def add_title(slide, text, left=Inches(0.6), top=Inches(0.45), width=Inches(12)):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.7))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, 28, bold=True, color=DARK)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left,
        top + Inches(0.58),
        Inches(1.2),
        Inches(0.08),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.fill.background()


def add_body(slide, lines, left=Inches(0.6), top=Inches(1.35), width=Inches(12), size=18):
    box = slide.shapes.add_textbox(left, top, width, Inches(5.3))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = line
        set_run(run, size, color=GRAY)


def add_card(slide, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = ORANGE_BORDER
    card.line.width = Pt(1.5)
    return card


def add_card_title(slide, text, left, top, width):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, 15, bold=True, color=ORANGE)


def add_card_body(slide, lines, left, top, width, height, size=15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = line
        set_run(run, size, color=GRAY)


def add_caption(slide, text, left, top, width):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.28))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    set_run(run, 12, bold=True, color=ORANGE)


def try_picture(slide, name, left, top, max_width, max_height):
    """等比縮放置入框內（不強制拉伸），置中對齊。"""
    path = SHOT / name
    if not path.exists():
        return False
    from PIL import Image

    with Image.open(path) as im:
        iw, ih = im.size
    if iw <= 0 or ih <= 0:
        return False

    max_w = int(max_width)
    max_h = int(max_height)
    scale = min(max_w / iw, max_h / ih)
    pic_w = max(1, int(iw * scale))
    pic_h = max(1, int(ih * scale))
    x = int(left) + (max_w - pic_w) // 2
    y = int(top) + (max_h - pic_h) // 2
    slide.shapes.add_picture(str(path), x, y, width=pic_w, height=pic_h)
    return True


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_pixel_bar(slide)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    page = 0

    # 1 封面
    page += 1
    s = new_slide(prs)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'NoteAdventure'
    set_run(run, 48, bold=True, color=DARK)

    box2 = s.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.5), Inches(0.5))
    p = box2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = '音樂練習 RPG 系統 ｜ 期中專題報告'
    set_run(run, 24, color=ORANGE)

    box3 = s.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.5), Inches(1.4))
    tf = box3.text_frame
    tf.word_wrap = True
    for i, line in enumerate([
        '從老師難以追蹤練習出發',
        '用遊戲化激勵上傳，並以實際畫面展示：Console → Streamlit → 學生 UI',
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        set_run(run, 18, color=GRAY)
    add_footer(s, page)

    # 2 大綱
    page += 1
    s = new_slide(prs)
    add_title(s, '報告大綱')
    add_body(s, [
        '1. 開發原點：想解決什麼問題？',
        '2. 解方方向：遊戲化練習與上傳動機',
        '3. 開發演進與實際操作畫面',
        '　　• Console 一小段',
        '　　• Streamlit 老師前端操作',
        '　　• 學生 HTML UI 操作',
        '4. 系統架構與重要功能邏輯',
        '5. 目前成果與下一步',
    ], size=20)
    add_footer(s, page)

    # 3 原點
    page += 1
    s = new_slide(prs)
    add_title(s, '開發原點：老師的痛點')
    add_card(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(5.0))
    add_body(s, [
        '一對一音樂課裡，老師最常遇到的不是「不會教」，而是：',
        '',
        '• 課與課之間，難以掌握學生有沒有真的練琴',
        '• 練習成果分散在訊息、檔案、口頭回報，難以追蹤',
        '• 缺少可持續累積的練習紀錄與回饋迴路',
        '',
        '因此專案原點是：讓老師更容易掌握與追蹤學生練習狀況。',
    ], left=Inches(0.95), top=Inches(1.7), width=Inches(11.3), size=20)
    add_footer(s, page)

    # 4 解方
    page += 1
    s = new_slide(prs)
    add_title(s, '解方方向：遊戲化練習')
    add_card(s, Inches(0.55), Inches(1.4), Inches(3.9), Inches(4.7))
    add_card_title(s, '學生端動機', Inches(0.8), Inches(1.65), Inches(3.4))
    add_card_body(s, [
        '• 上傳練習錄音',
        '• 完成任務拿 XP／金幣',
        '• 寵物成長與商店獎勵',
        '',
        '希望學生因為「想上傳」',
        '而增加彈琴次數',
    ], Inches(0.8), Inches(2.2), Inches(3.4), Inches(3.5), size=16)

    add_card(s, Inches(4.7), Inches(1.4), Inches(3.9), Inches(4.7))
    add_card_title(s, '老師端價值', Inches(4.95), Inches(1.65), Inches(3.4))
    add_card_body(s, [
        '• 建立學生與任務',
        '• 指派任務並看進度',
        '• 查看上傳紀錄',
        '• 管理商店與兌換',
        '',
        '把零散練習變成',
        '可追蹤的資料流',
    ], Inches(4.95), Inches(2.2), Inches(3.4), Inches(3.5), size=16)

    add_card(s, Inches(8.85), Inches(1.4), Inches(3.9), Inches(4.7))
    add_card_title(s, '核心迴路', Inches(9.1), Inches(1.65), Inches(3.4))
    add_card_body(s, [
        '指派任務',
        '→ 上傳練習',
        '→ 更新進度',
        '→ 發放獎勵',
        '→ 寵物成長',
        '',
        '用獎勵把練習',
        '變成可持續行為',
    ], Inches(9.1), Inches(2.2), Inches(3.4), Inches(3.5), size=16)
    add_footer(s, page)

    # 5 三階段
    page += 1
    s = new_slide(prs)
    add_title(s, '開發演進：三個階段')
    steps = [
        ('① Console', '先把規則與流程寫清楚\n確認業務邏輯正確'),
        ('② Streamlit', '用同一套邏輯接網頁\n讓老師能實際操作'),
        ('③ 學生專用 UI', '補上溫度與沉浸感\n提高學生使用意願'),
    ]
    for i, (title, body) in enumerate(steps):
        left = Inches(0.6 + i * 4.2)
        add_card(s, left, Inches(1.6), Inches(3.9), Inches(4.4))
        add_card_title(s, title, left + Inches(0.25), Inches(1.9), Inches(3.4))
        add_card_body(
            s,
            body.split('\n'),
            left + Inches(0.25),
            Inches(2.6),
            Inches(3.4),
            Inches(3.0),
            size=18,
        )
    add_footer(s, page)

    # 6 Console 操作畫面（一小段）
    page += 1
    s = new_slide(prs)
    add_title(s, '階段一：Console 操作畫面（一小段）')
    add_body(s, [
        '先用文字選單把流程跑通；以下只裁選單重點，方便看清楚。',
    ], top=Inches(1.15), width=Inches(12), size=16)
    pics = [
        ('c01_console_main.png', '主選單'),
        ('c02_console_teacher.png', '老師選單'),
        ('c03_console_student.png', '學生選單'),
    ]
    for i, (name, label) in enumerate(pics):
        left = Inches(0.35 + i * 4.35)
        add_card(s, left, Inches(1.75), Inches(4.15), Inches(5.0))
        add_caption(s, label, left + Inches(0.15), Inches(1.85), Inches(3.85))
        try_picture(s, name, left + Inches(0.15), Inches(2.2), Inches(3.85), Inches(4.35))
    add_footer(s, page)

    # 7 Streamlit 老師操作 1
    page += 1
    s = new_slide(prs)
    add_title(s, '階段二：Streamlit 老師前端操作')
    add_body(s, [
        '同一套 Flask API；畫面已裁掉側欄，放大首頁統計與學生列表。',
    ], top=Inches(1.15), width=Inches(12), size=16)
    add_card(s, Inches(0.4), Inches(1.7), Inches(6.2), Inches(5.1))
    add_caption(s, '老師首頁 Dashboard', Inches(0.55), Inches(1.8), Inches(5.9))
    try_picture(s, '02_teacher_home.png', Inches(0.55), Inches(2.15), Inches(5.9), Inches(4.45))

    add_card(s, Inches(6.75), Inches(1.7), Inches(6.2), Inches(5.1))
    add_caption(s, '學生列表', Inches(6.9), Inches(1.8), Inches(5.9))
    try_picture(s, '03_teacher_students.png', Inches(6.9), Inches(2.15), Inches(5.9), Inches(4.45))
    add_footer(s, page)

    # 8 Streamlit 老師操作 2
    page += 1
    s = new_slide(prs)
    add_title(s, 'Streamlit：任務／上傳／商店')
    tiles = [
        ('04_teacher_quests.png', '任務清單'),
        ('05_teacher_uploads.png', '上傳紀錄'),
        ('06_teacher_shop.png', '商品清單'),
        ('07_teacher_redeem.png', '待兌換商品'),
    ]
    for i, (name, label) in enumerate(tiles):
        left = Inches(0.35 + (i % 2) * 6.5)
        top = Inches(1.2 + (i // 2) * 2.95)
        add_card(s, left, top, Inches(6.3), Inches(2.8))
        add_caption(s, label, left + Inches(0.12), top + Inches(0.06), Inches(6.0))
        try_picture(s, name, left + Inches(0.12), top + Inches(0.35), Inches(6.05), Inches(2.3))
    add_footer(s, page)

    # 9 轉折
    page += 1
    s = new_slide(prs)
    add_title(s, '關鍵轉折：操作頁面溫度不足')
    add_card(s, Inches(0.5), Inches(1.35), Inches(6.2), Inches(5.2))
    try_picture(s, '09_student_home.png', Inches(0.65), Inches(1.5), Inches(5.9), Inches(4.9))
    add_body(s, [
        '左圖是早期 Streamlit 學生頁：',
        '功能齊全，但偏工具感。',
        '',
        '若學生不想開，上傳就變少；',
        '上傳變少，老師又追不到練習。',
        '',
        '下一步：做學生專用溫暖 UI，',
        '讓遊戲化真的被使用。',
    ], left=Inches(7.0), top=Inches(1.7), width=Inches(5.6), size=18)
    add_footer(s, page)

    # 10 學生 UI 主頁
    page += 1
    s = new_slide(prs)
    add_title(s, '階段三：學生專用 UI（主頁地圖）')
    add_body(s, [
        '地圖式主頁：琴房、任務、商店、寵物、教室（裁切外框雜訊，放大地圖本體）。',
    ], top=Inches(1.15), width=Inches(12), size=16)
    add_card(s, Inches(0.9), Inches(1.65), Inches(11.5), Inches(5.2))
    try_picture(s, 'html_02_home.png', Inches(1.05), Inches(1.8), Inches(11.2), Inches(4.9))
    add_footer(s, page)

    # 11 學生 UI 操作畫面（直向四欄，避免扁壓變形）
    page += 1
    s = new_slide(prs)
    add_title(s, '學生 UI 實際操作畫面')
    tiles = [
        ('html_03_practice.png', '琴房｜上傳'),
        ('html_04_quests.png', '任務欄'),
        ('html_05_pet.png', '寵物'),
        ('html_06_shop.png', '商店'),
    ]
    for i, (name, label) in enumerate(tiles):
        left = Inches(0.28 + i * 3.28)
        add_card(s, left, Inches(1.25), Inches(3.15), Inches(5.5))
        add_caption(s, label, left + Inches(0.1), Inches(1.35), Inches(2.95))
        try_picture(s, name, left + Inches(0.1), Inches(1.7), Inches(2.95), Inches(4.85))
    add_footer(s, page)

    # 12 教室看板
    page += 1
    s = new_slide(prs)
    add_title(s, '學生 UI：教室看板（兩廳院節目）')
    add_body(s, [
        '重點裁切看板區：近一個月音樂節目海報＋資訊自動輪播。',
    ], top=Inches(1.15), width=Inches(12), size=16)
    add_card(s, Inches(0.9), Inches(1.65), Inches(11.5), Inches(5.2))
    try_picture(s, 'html_07_class.png', Inches(1.05), Inches(1.8), Inches(11.2), Inches(4.9))
    add_footer(s, page)

    # 13 架構
    page += 1
    s = new_slide(prs)
    add_title(s, '目前架構：一套邏輯，兩套介面')
    add_body(s, [
        '學生 HTML／JS  ──┐',
        '                  ├──→  Flask API  ──→  MySQL',
        '老師 Streamlit ──┘         │',
        '                            └── services（業務邏輯）',
        '',
        '• Console 時期長出的服務層，成為 API 的基礎',
        '• 老師端偏管理效率；學生端偏使用溫度',
        '• 兩邊共用登入、任務、上傳、獎勵等規則',
    ], size=20)
    add_footer(s, page)

    # 14 核心邏輯
    page += 1
    s = new_slide(prs)
    add_title(s, '重要功能邏輯（精簡）')
    add_card(s, Inches(0.5), Inches(1.35), Inches(6.0), Inches(5.1))
    add_card_title(s, '練習上傳 → 任務進度', Inches(0.75), Inches(1.55), Inches(5.5))
    add_card_body(s, [
        '1. 學生上傳練習錄音',
        '2. 寫入 practice_upload',
        '3. 依任務條件更新 student_quest',
        '4. 達標後完成任務',
        '5. 發 XP／金幣（避免重複發獎）',
        '6. 依累積 XP 更新寵物階段',
    ], Inches(0.75), Inches(2.15), Inches(5.5), Inches(3.9), size=16)

    add_card(s, Inches(6.8), Inches(1.35), Inches(5.9), Inches(5.1))
    add_card_title(s, '其他關鍵能力', Inches(7.05), Inches(1.55), Inches(5.4))
    add_card_body(s, [
        '• 老師建立／指派任務',
        '• 商店購買與兌換',
        '• Dashboard 掌握整體狀態',
        '• 教室看板：兩廳院音樂節目',
        '',
        '原則：先保證規則正確，',
        '再用介面提高使用率。',
    ], Inches(7.05), Inches(2.15), Inches(5.4), Inches(3.9), size=16)
    add_footer(s, page)

    # 15 進度
    page += 1
    s = new_slide(prs)
    add_title(s, '目前進度')
    add_body(s, [
        '已完成：',
        '• Console 核心流程與資料庫規則',
        '• Flask API（登入、任務、上傳、商店、寵物、Dashboard、教室看板）',
        '• Streamlit 老師端操作介面',
        '• 學生 HTML UI（主頁、琴房、任務、商店、寵物、設定、教室）',
        '',
        '持續優化：',
        '• 背包完整頁、更多任務條件、使用體驗細節',
    ], size=18)
    add_footer(s, page)

    # 16 結語
    page += 1
    s = new_slide(prs)
    add_title(s, '結語')
    add_body(s, [
        'NoteAdventure 不是先做畫面，而是：',
        '',
        '痛點 → 遊戲化解方 → 邏輯驗證（Console）',
        '→ 管理介面（Streamlit）→ 學生溫度 UI',
        '',
        '希望最後同時做到：',
        '老師追得到練習，學生也願意持續上傳練琴。',
        '',
        '謝謝聆聽，歡迎提問。',
    ], size=20)
    add_footer(s, page)

    prs.save(OUT)
    print(f'已輸出：{OUT}')
    print(f'截圖來源：{SHOT}')


if __name__ == '__main__':
    build()
